import time, traceback
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from zoneinfo import ZoneInfo

from db_utils import save_dataframe, get_engine
from config import POSTGRES_SCHEMA


def log_message(module: str, message: str):
    timestamp = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] [{module}] {message}", flush=True)


# ----------------------------------------------------------------------
# 1) Carregar Excel
# ----------------------------------------------------------------------
def carregar_excel(excel_file, sheet_name, table_name, drop_key = None):
    log_message("api_businessUnits", f"Lendo Excel: {excel_file} | aba={sheet_name} | tabela={table_name}")

    wb = load_workbook(excel_file, data_only=True)
    ws = wb[sheet_name]

    # Localiza a tabela nomeada
    table = ws.tables[table_name]
    ref = table.ref
    col_start = ''.join([c for c in ref.split(":")[0] if c.isalpha()])
    col_end   = ''.join([c for c in ref.split(":")[1] if c.isalpha()])
    usecols = f"{col_start}:{col_end}"
    start_row = int(''.join([c for c in ref.split(":")[0] if c.isdigit()]))
    df = pd.read_excel(
        excel_file,
        sheet_name=sheet_name,
        usecols=usecols,
        header=0,
        skiprows=start_row - 1)
    # Remove linhas sem chave
    if drop_key and drop_key in df.columns:
        df = df.dropna(subset=[drop_key])

    log_message("api_businessUnits", f"Excel carregado com {len(df)} linhas.")
    return df
# ----------------------------------------------------------------------
# 2) Salvar no banco
# ----------------------------------------------------------------------
def save_table(engine, df, table_name, module="api_businessUnits"):
    start = time.time()
    started_at = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")

    try:
        if isinstance(df, pd.DataFrame) and not df.empty:
            save_dataframe(df, table_name, POSTGRES_SCHEMA)
            elapsed = time.time() - start
            log_message(module, f"{table_name} - status: ✅ ok - início: {started_at} - tempo: {elapsed:.2f}s - linhas: {len(df)}")
            return {"table": table_name, "status": "✅ ok", "rows": len(df), "elapsed": elapsed}

        else:
            elapsed = time.time() - start
            log_message(module, f"{table_name} - status: ⚠️ vazio - início: {started_at} - tempo: {elapsed:.2f}s - linhas: 0")
            return {"table": table_name, "status": "⚠️ vazio", "rows": 0, "elapsed": elapsed}

    except Exception as e:
        elapsed = time.time() - start
        log_message(module, f"❌ Erro ao salvar {table_name}: {e}")
        log_message(module, traceback.format_exc())
        return {"table": table_name, "status": "❌ erro", "rows": None, "elapsed": elapsed}


# ----------------------------------------------------------------------
# 3) Main
# ----------------------------------------------------------------------
def main():
    global_start = time.time()
    started_at = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")

    try:
        excel_file = "/scripts/Logs.xlsx"
        df_bu = carregar_excel(
            excel_file,
            sheet_name="billId_documentNumber",
            table_name="tb_BU",
            drop_key="bill_doc_number")

        df_dacao = carregar_excel(
            excel_file,
            sheet_name="Dacao_REC",
            table_name="undDacao",
            drop_key="mainUnit_CompanyId")  

        # 2) Salvar no banco
        engine = get_engine()
        resumo_bu = save_table(engine, df_bu, "BU_businessUnits")
        resumo_dacao = save_table(engine, df_dacao, "DacaoREC")
        elapsed = time.time() - global_start

        log_message("api_businessUnits", "Resumo da execução:")
        for r in [resumo_bu, resumo_dacao]:
            log_message(
                "api_businessUnits",
                f"{r['table']} - status: {r['status']} - início: {started_at} - tempo total: {elapsed:.2f}s - linhas: {r['rows']}")

    except Exception as e:
        elapsed = time.time() - global_start
        log_message("api_businessUnits", f"❌ Erro ao processar após {elapsed:.2f}s: {e}")
        log_message("api_businessUnits", traceback.format_exc())

if __name__ == "__main__":
    main()
