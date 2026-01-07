import time, traceback
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from zoneinfo import ZoneInfo

from db_utils import save_dataframe, get_engine
from config import POSTGRES_SCHEMA


def log_message(module: str, message: str):
    timestamp = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] [{module}] {message}", flush=True)


# ----------------------------------------------------------------------
# 1) Carregar Excel
# ----------------------------------------------------------------------
def carregar_excel(excel_file, sheet_name, table_name):
    log_message("api_businessUnits", f"Lendo Excel: {excel_file} | aba={sheet_name} | tabela={table_name}")

    wb = load_workbook(excel_file, data_only=True)
    ws = wb[sheet_name]

    # Localiza a tabela nomeada
    table = ws.tables[table_name]
    ref = table.ref
    col_start = ''.join([c for c in ref.split(":")[0] if c.isalpha()])
    col_end   = ''.join([c for c in ref.split(":")[1] if c.isalpha()])
    usecols = f"{col_start}:{col_end}"
    start_row = int(''.join([c for c in ref.split(":")[0] if c.isdigit()]))
    df = pd.read_excel(
        excel_file,
        sheet_name=sheet_name,
        usecols=usecols,
        header=0,
        skiprows=start_row - 1)
    # Remove linhas sem chave
    df = df.dropna(subset=["bill_doc_number"])

    log_message("api_businessUnits", f"Excel carregado com {len(df)} linhas.")
    return df
# ----------------------------------------------------------------------
# 2) Salvar no banco
# ----------------------------------------------------------------------
def save_business_units(engine, df, module="api_businessUnits"):
    start = time.time()
    started_at = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")

    try:
        if isinstance(df, pd.DataFrame) and not df.empty:
            save_dataframe(df, "BU_businessUnits", POSTGRES_SCHEMA)
            elapsed = time.time() - start
            log_message(module, f"businessUnits - status: ✅ ok - início: {started_at} - tempo: {elapsed:.2f}s - linhas: {len(df)}")
            return {"table": "businessUnits", "status": "✅ ok", "rows": len(df), "elapsed": elapsed}

        else:
            elapsed = time.time() - start
            log_message(module, f"businessUnits - status: ⚠️ vazio - início: {started_at} - tempo: {elapsed:.2f}s - linhas: 0")
            return {"table": "businessUnits", "status": "⚠️ vazio", "rows": 0, "elapsed": elapsed}

    except Exception as e:
        elapsed = time.time() - start
        log_message(module, f"❌ Erro ao salvar businessUnits: {e}")
        log_message(module, traceback.format_exc())
        return {"table": "businessUnits", "status": "❌ erro", "rows": None, "elapsed": elapsed}


# ----------------------------------------------------------------------
# 3) Main
# ----------------------------------------------------------------------
def main():
    global_start = time.time()
    started_at = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")

    try:
        excel_file = "/scripts/Logs.xlsx"
        sheet_name = "billId_documentNumber"
        table_name = "tb_BU"

        # 1) Carregar Excel
        df = carregar_excel(excel_file, sheet_name, table_name)
        # 2) Salvar no banco
        engine = get_engine()
        resumo = save_business_units(engine, df)
        # 3) Resumo final
        elapsed = time.time() - global_start
        status_global = resumo["status"]
        linhas_total = resumo["rows"]

    except Exception as e:
        elapsed = time.time() - global_start
        log_message("api_businessUnits", f"❌ Erro ao processar após {elapsed:.2f}s: {e}")
        log_message("api_businessUnits", traceback.format_exc())


if __name__ == "__main__":
    main()
