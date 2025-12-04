import logging, time, traceback
import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import inspect, text
from db_utils import get_engine
from datetime import datetime
from zoneinfo import ZoneInfo

def log_message(module: str, message: str):
    timestamp = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] [{module}] {message}")

def carregar_excel(excel_file, sheet_name, table_name):
    """Lê uma tabela nomeada de um Excel e retorna um DataFrame."""
    wb = load_workbook(excel_file, data_only=True)
    ws = wb[sheet_name]

    table = ws.tables[table_name]
    ref = table.ref
    cols_range = ref.split(":")
    col_start = ''.join([c for c in cols_range[0] if c.isalpha()])
    col_end   = ''.join([c for c in cols_range[1] if c.isalpha()])
    usecols = f"{col_start}:{col_end}"
    start_row = int(''.join([c for c in cols_range[0] if c.isdigit()]))

    df = pd.read_excel(excel_file, sheet_name=sheet_name, usecols=usecols, header=0, skiprows=start_row-1)
    df = df.dropna(subset=["bill_doc_number"])
    return df

def atualizar_tabela(engine, nome, df_on):
    """Atualiza uma tabela no banco com base em um DataFrame."""
    start_time = datetime.now(ZoneInfo("America/Sao_Paulo"))
    start = time.time()
    try:
        with engine.begin() as conn:
            inspector = inspect(conn)
            if not inspector.has_table(nome):
                df_on.head(0).to_sql(nome, conn, if_exists='replace', index=False)
            conn.execute(text(f'TRUNCATE TABLE "{nome}" RESTART IDENTITY CASCADE'))
            df_on.to_sql(nome, conn, if_exists='append', index=False)

        return {"table": nome, "status": "✅ ok", "rows": len(df_on)}
    except Exception as e:
        elapsed = time.time() - start
        log_message("sienge_excel_bu", f"Erro ao atualizar a tabela {nome}: {e}")
        log_message("sienge_excel_bu", traceback.format_exc())
        return {"table": nome, "status": "❌ erro", "rows": None}


def main():
    global_start = time.time()
    started_at = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%d %H:%M:%S")

    # Caminho e tabela específicos deste Excel
    excel_file = "/scripts/Logs.xlsx"
    sheet_name = "billId_documentNumber"
    table_name = "tb_BU"

    # 1. Ler dados do Excel
    df = carregar_excel(excel_file, sheet_name, table_name)
    engine = get_engine()
    resumo = atualizar_tabela(engine, "BU_businessUnits", df)

    # --- Resumo final ---
    elapsed = time.time() - global_start
    log_message("sienge_excel_bu", "Resumo da execução:")
    log_message(
        "sienge_excel_bu",
        f"{resumo['table']} - status: {resumo['status']} - início: {started_at} - tempo total: {elapsed:.2f}s - linhas: {resumo['rows']}")

if __name__ == "__main__":
    main()